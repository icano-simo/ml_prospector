"""Los handles de Instagram que no sirven, listos para corregir a mano.

CERO peticiones: solo lee `ig_signals.csv`, los crudos de `ig_raw/` y los
registros de Model Match.

Tres grupos, y conviene no mezclarlos porque se arreglan distinto:

    no_encontrado   el handle que encontro la busqueda NO EXISTE en Instagram.
                    Alguien se equivoco, o la cuenta se borro o se renombro.
    handle_dudoso   existe, pero el perfil no parece de esa persona.
    confianza baja  existe y se leyo, pero la verificacion no paso: el nombre
                    del perfil no se parece y no hay señal inmobiliaria. Sus
                    señales NO se usan (ver la skill instagram-busqueda).

La columna que hay que llenar es `Handle correcto`. Vale tambien `NO TIENE`,
que es una respuesta: hace que no se le vuelva a buscar.

Uso:
    python realtor_scraper/handles_para_revisar.py
"""
from __future__ import annotations

import csv
import glob
import json
import os
import sys

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

CSV_IG = os.path.join(RAIZ, "realtor_scraper", "output", "ig_signals.csv")
DIR_CRUDO = os.path.join(RAIZ, "realtor_scraper", "output", "ig_raw")
DIR_MM = os.path.join(RAIZ, "data", "trabajo", "mm_por_realtor")
SALIDA = os.path.join(RAIZ, "data", "salida")
ARCHIVO = os.path.join(SALIDA, "revisar_handles_instagram.xlsx")

#: Que se escribe en la columna «Por que hay que revisarlo». Es el motivo, no
#: el estado crudo: el estado dice `no_encontrado` y eso no le dice a nadie
#: que tiene que hacer.
MOTIVOS = {
    "no_encontrado": "el handle NO EXISTE en Instagram: la cuenta se borró, "
                     "se renombró, o la búsqueda se equivocó",
    "handle_dudoso": "la cuenta existe pero no parece de esta persona",
}

#: ⚠ `sin_grid` NO entra por si solo, y la distincion importa.
#:
#: Son 114 perfiles que cargaron sin cuadricula de posts. Eso NO es un handle
#: equivocado: la mayoria tiene confianza alta o media, o sea que el handle
#: SI es de esa persona. Lo que fallo es la lectura del contenido, y eso se
#: arregla con `--reintentar-ilegibles`, no buscando otro handle.
#:
#: Meterlos aqui mandaria a alguien a buscar a mano 114 cuentas que ya estan
#: bien identificadas. Solo entran los que ademas no verifican.
SOLO_SI_NO_VERIFICA = {"sin_grid"}
MOTIVO_BAJA = ("la cuenta existe y se leyó, pero NO verifica: el nombre del "
               "perfil no se parece al del realtor y no hay señal "
               "inmobiliaria. Sus señales de contenido no se usan")


def cargar_mm() -> dict:
    """Lo que Model Match sabe de cada uno, por correo."""
    por_correo = {}
    for a in glob.glob(os.path.join(DIR_MM, "*.json")):
        with open(a, encoding="utf-8") as fh:
            f = json.load(fh)
        for c in (f.get("emails_mmi") or []):
            if c:
                por_correo[c.strip().lower()] = f
    return por_correo


def cargar_crudos() -> dict:
    """El diagnostico de la peticion, por correo del objetivo.

    Los crudos se nombran por `clave`, no por handle, asi que el cruce va por
    el correo que quedo guardado dentro, en `objetivo`.
    """
    por_correo = {}
    for a in glob.glob(os.path.join(DIR_CRUDO, "*.json")):
        try:
            with open(a, encoding="utf-8") as fh:
                d = json.load(fh)
        except (json.JSONDecodeError, OSError):
            continue
        c = ((d.get("objetivo") or {}).get("email") or "").strip().lower()
        if c:
            por_correo[c] = d
    return por_correo


def escribir(ws, titulos, anchos, filas):
    ws.append(titulos)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(vertical="top", wrap_text=True)
    for f in filas:
        ws.append(f)
    for i, an in enumerate(anchos, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = an
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main() -> None:
    csv.field_size_limit(10_000_000)
    with open(CSV_IG, encoding="utf-8-sig", newline="") as fh:
        filas_csv = list(csv.DictReader(fh))
    mm = cargar_mm()
    crudos = cargar_crudos()

    filas, cuenta = [], {}
    for f in filas_csv:
        estado = (f.get("estado_perfil") or "").strip()
        conf = (f.get("handle_confianza") or "").strip()
        if estado in MOTIVOS:
            grupo, motivo = estado, MOTIVOS[estado]
        elif conf == "baja":
            grupo, motivo = "confianza_baja", MOTIVO_BAJA
            if estado in SOLO_SI_NO_VERIFICA:
                motivo = ("la cuenta cargó sin cuadrícula de posts Y además "
                          "no verifica, así que puede ser el handle y no la "
                          "lectura")
        else:
            continue
        cuenta[grupo] = cuenta.get(grupo, 0) + 1

        correo = (f.get("email") or "").strip().lower()
        reg = mm.get(correo) or {}
        crudo = crudos.get(correo) or {}
        handle = (f.get("handle") or "").strip().lstrip("@")
        filas.append([
            "",                                   # Handle correcto
            f.get("nombre"), f.get("estado"),
            reg.get("mm_ciudad") or "",
            reg.get("mm_brokerage") or reg.get("brokerage_mmi") or "",
            correo,
            " · ".join(reg.get("telefonos_mmi") or []),
            handle,
            "https://instagram.com/%s" % handle if handle else "",
            grupo, motivo,
            crudo.get("http_status") or "",
            " · ".join(crudo.get("errores") or [])[:200],
            "SÍ" if reg.get("mm_id") else "no",
            reg.get("realtor_id") or "",
        ])

    if not filas:
        print("no hay handles que revisar")
        return

    # Primero los que ya tienen Model Match: arreglar su handle vale el doble,
    # porque su ficha ya esta pagada y lo unico que falta es el contenido.
    filas.sort(key=lambda r: (r[13] != "SÍ", r[9], r[1] or ""))

    wb = Workbook()
    ws = wb.active
    ws.title = "Handles a corregir"
    escribir(ws,
             ["Handle correcto", "Realtor", "Estado", "Ciudad (MM)",
              "Brokerage", "Correo", "Teléfonos", "Handle que se probó",
              "Abrir en Instagram", "Grupo", "Por qué hay que revisarlo",
              "HTTP", "Errores de la petición", "¿Tiene Model Match?",
              "realtor_id"],
             [22, 26, 8, 18, 28, 32, 24, 24, 38, 16, 62, 7, 40, 16, 36],
             filas)
    amarillo = PatternFill("solid", fgColor="FFF3CD")
    for fila in ws.iter_rows(min_row=2):
        fila[0].fill = amarillo
        fila[10].alignment = Alignment(wrap_text=True, vertical="top")

    ws2 = wb.create_sheet("Cómo se usa")
    con_mm = sum(1 for r in filas if r[13] == "SÍ")
    notas = [
        ["Qué es esto", ""],
        ["", "Los %d realtors cuyo handle de Instagram NO sirve. Todos tienen "
             "un handle anotado: lo que falla es que no existe, que no es de "
             "esa persona, o que no se pudo verificar." % len(filas)],
        ["", "%d de ellos YA tienen su ficha de Model Match comprada y "
             "pagada. En ésos, corregir el handle es lo único que falta para "
             "tener las dos capas completas, así que van primero en la lista."
         % con_mm],
        ["Qué hay que hacer", ""],
        ["", "En la columna A escribir el handle correcto, sin la @. Si la "
             "persona no tiene Instagram, escribir NO TIENE: eso también es "
             "una respuesta y evita que se le vuelva a buscar."],
        ["", "La columna «Abrir en Instagram» lleva al handle que se probó, "
             "para ver con los propios ojos qué hay ahí."],
        ["Los tres grupos, que se arreglan distinto", ""],
        ["no_encontrado",
         "El handle no existe. Suele ser una cuenta borrada o renombrada, o "
         "un error de la búsqueda. Es el grupo donde más se gana buscando a "
         "mano, porque la persona probablemente sí tiene cuenta."],
        ["handle_dudoso",
         "La cuenta existe pero no parece de esta persona. Hay que confirmar "
         "si es un homónimo o si es ella con un perfil que no lo declara."],
        ["confianza_baja",
         "La cuenta existe y se leyó, pero no verifica: el nombre del perfil "
         "no se parece y no hay señal inmobiliaria. Puede ser su cuenta "
         "PERSONAL —y entonces el handle está bien pero no sirve para esto— o "
         "puede ser otra persona."],
        ["Por qué el sistema no lo resolvió solo", ""],
        ["", "La regla pide dos cosas: que el nombre del perfil se parezca al "
             "del realtor, y que haya señal inmobiliaria en la bio o en los "
             "posts. Si no se cumplen, el handle se guarda con confianza baja "
             "y sus señales NO se usan."],
        ["", "La razón: un handle equivocado no produce un error, produce "
             "cuarenta señales de contenido sobre la persona equivocada, con "
             "el mismo aspecto que las correctas."],
        ["Cuidado con este archivo", ""],
        ["", "Lleva nombres, correos y teléfonos de personas reales. No se "
             "versiona y no se comparte fuera del equipo."],
    ]
    escribir(ws2, ["Concepto", "Qué significa"], [20, 108], notas)
    for fila in ws2.iter_rows(min_row=2, max_col=2):
        fila[1].alignment = Alignment(wrap_text=True, vertical="top")
        if fila[0].value and not fila[1].value:
            fila[0].font = Font(bold=True, size=11)

    os.makedirs(SALIDA, exist_ok=True)
    wb.save(ARCHIVO)
    print("handles a revisar: %d" % len(filas))
    for k, v in sorted(cuenta.items(), key=lambda x: -x[1]):
        print("   %-16s %d" % (k, v))
    print("")
    print("con Model Match ya pagado: %d  (van primero)" % con_mm)
    print("guardado en %s" % ARCHIVO)


if __name__ == "__main__":
    main()
