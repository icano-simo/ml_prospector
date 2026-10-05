"""Los perfiles de Instagram ya raspados, a Excel.

Lee `output/ig_signals.csv`, que es lo que produce `--parsear`, y lo deja en
`data/salida/instagram_perfiles.xlsx`.

**Avisa si el CSV esta atrasado respecto del crudo.** El scraper guarda un
JSON por perfil en `output/ig_raw/` y el CSV solo se actualiza cuando se corre
`--parsear`. El 30/09 habia 513 crudos y 300 filas en el CSV: 213 perfiles
raspados que no estaban en ningun lado. Un exportador que no mire eso entrega
un Excel incompleto con cara de completo.

Uso:
    python realtor_scraper/exportar_ig_excel.py
"""
from __future__ import annotations

import csv
import glob
import os
import sys

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(RAIZ, "realtor_scraper", "output")
SALIDA = os.path.join(RAIZ, "data", "salida")
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

#: Excel corta cualquier celda a 32.767 caracteres. Las columnas de texto
#: crudo pasan eso, y si se escriben enteras el archivo no abre.
TOPE_CELDA = 32000


def main() -> None:
    ruta_csv = os.path.join(OUT, "ig_signals.csv")
    if not os.path.exists(ruta_csv):
        raise SystemExit(
            "no existe %s\nCorre primero, desde realtor_scraper/:\n"
            "    python -m instagram.finder --parsear" % ruta_csv)

    with open(ruta_csv, encoding="utf-8-sig", newline="") as fh:
        filas = list(csv.DictReader(fh))
    if not filas:
        raise SystemExit("el CSV esta vacio")

    crudos = len(glob.glob(os.path.join(OUT, "ig_raw", "*.json")))
    print("perfiles en el CSV : %d" % len(filas))
    print("crudos en ig_raw   : %d" % crudos)
    atraso = crudos - len(filas)
    if atraso > 0:
        print("")
        print("⚠  HAY %d PERFILES RASPADOS QUE NO ESTAN EN EL CSV." % atraso)
        print("   El Excel saldra con %d y no con %d. Para incluirlos, antes:"
              % (len(filas), crudos))
        print("       cd realtor_scraper")
        print("       python -m instagram.finder --parsear")
        print("")

    cols = list(filas[0])
    wb = Workbook()
    ws = wb.active
    ws.title = "Perfiles IG"
    ws.append(cols)
    for f in filas:
        ws.append([(f.get(c) or "")[:TOPE_CELDA] for c in cols])

    for i, c in enumerate(cols, 1):
        ws.cell(row=1, column=i).font = Font(bold=True, color="FFFFFF", size=10)
        ws.cell(row=1, column=i).fill = PatternFill("solid", fgColor="1F3864")
        ws.cell(row=1, column=i).alignment = Alignment(wrap_text=True,
                                                       vertical="center")
        # Las de texto largo, angostas: si no, empujan todo fuera de pantalla.
        ancho = 60 if c in ("captions_texto", "comentarios_texto",
                            "citas_por_etiqueta") else 18
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions

    os.makedirs(SALIDA, exist_ok=True)
    ruta = os.path.join(SALIDA, "instagram_perfiles.xlsx")
    try:
        wb.save(ruta)
    except PermissionError:
        import datetime as dt
        ruta = os.path.join(SALIDA, "instagram_perfiles_%s.xlsx"
                            % dt.datetime.now().strftime("%H%M"))
        wb.save(ruta)
        print("⚠ el archivo estaba abierto en Excel; se guardo al lado")

    print("%d perfiles x %d columnas" % (len(filas), len(cols)))
    print("guardado en %s" % ruta)


if __name__ == "__main__":
    main()
