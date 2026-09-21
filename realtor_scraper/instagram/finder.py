"""Orquestador de la capa de Instagram.

Que cambio respecto de la version anterior
------------------------------------------
La version anterior devolvia catorce booleanos y un `{}` cuando algo fallaba, y
ese `{}` se volvia `False` en las catorce columnas. Esta devuelve un objeto con
**estado del perfil**, **confianza del handle** y una **mascara de
disponibilidad por señal**.

La regla que ordena todo el modulo: *nunca se devuelve False donde corresponde
null.* Una señal ausente lleva su motivo, y el motivo importa -- un perfil
privado es un dato sobre la persona, un handle equivocado es un dato sobre
nuestro scraper.

Y la segunda: *si el handle no verifica, sus señales no se usan.* Un handle
equivocado no producia un error: producia catorce señales sobre la persona
equivocada.
"""
from __future__ import annotations

import random
import re
import time
import urllib.parse
from dataclasses import asdict, dataclass, field

from loguru import logger

from instagram import extraccion
from instagram.comentarios import (
    PerfilDeAudiencia,
    perfilar_audiencia,
    verificar_anonimato,
)
from instagram.estado import EstadoPerfil
from instagram.extraccion import PerfilCrudo
from instagram.idioma import PerfilDeIdioma, declara_acompanamiento, perfilar
from instagram.posts import (
    Cadencia,
    EngagementReal,
    RedDeMenciones,
    calcular_cadencia,
    calcular_engagement,
    construir_red,
    geotags_agregados,
    programas_agregados,
)
from instagram.verificacion import (
    Confianza,
    Verificacion,
    extraer_licencias_de_bio,
    verificar,
)

DDG = "https://duckduckgo.com/"

#: Candidatos a evaluar por realtor. La version anterior tomaba el primero que
#: apareciera en el HTML y por eso "encontraba" el 98,1%.
MAX_CANDIDATOS = 5

PAUSA_BUSQUEDA = (2.5, 5.0)

#: Rutas de Instagram que no son handles de persona.
_NO_SON_HANDLES = {
    "p", "reel", "reels", "explore", "accounts", "stories", "tv", "about",
    "privacy", "legal", "help", "tags", "directory", "developer", "api",
    "instagram", "web", "direct", "challenge", "emails", "session",
    "oauth", "graphql", "static", "images", "ajax", "your_activity",
}


# ── Resultado ─────────────────────────────────────────────────────────────────

@dataclass
class SenalesInstagram:
    """Lo que la capa de Instagram entrega al motor de reglas PACS.

    Cada campo que puede faltar es `None`, y `disponibilidad` dice por que.
    """

    # Identidad y estado
    handle: str | None = None
    url: str | None = None
    estado_perfil: str = EstadoPerfil.SIN_HANDLE.value
    motivo_estado: str = ""
    handle_confidence: str = Confianza.BAJA.value
    razon_confianza: str = ""
    candidatos_evaluados: list[dict] = field(default_factory=list)

    # Huella digital (S1)
    nombre_perfil: str | None = None
    bio: str | None = None
    seguidores: int | None = None
    n_posts_declarados: int | None = None
    categoria_declarada: str | None = None

    # Madurez tecnologica (S4)
    link_de_bio: str | None = None
    destino_link_de_bio: str | None = None
    titulos_de_destacadas: list[str] = field(default_factory=list)

    # Identidad cruzada
    licencias_en_bio: list[str] = field(default_factory=list)

    # Contenido (S3, S9) — CRUDO, no solo flags
    posts_crudos: list[dict] = field(default_factory=list)

    # Idioma (P-Q14)
    idioma: dict | None = None
    intensidad_pq14: int | None = None
    grado_pq14: str | None = None
    regla_pq14: str | None = None
    evidencia_ancla_idioma: str | None = None

    # Programas (P-Q01, P-Q07, P-Q17, P-Q19)
    programas: dict = field(default_factory=dict)

    # Relaciones (S6) y co-marketing (P-Q21)
    red_de_menciones: dict | None = None
    candidatas_a_lender: list[str] = field(default_factory=list)
    n_posts_co_marketing: int = 0

    # Produccion y cadencia (P-Q10, J-Q04)
    cadencia: dict | None = None
    n_posts_con_produccion: int = 0

    # R1 verdadero
    engagement: dict | None = None

    # S7 sub-estatal
    geotags: dict = field(default_factory=dict)

    # S8 (perfil de audiencia, agregado y anonimo)
    audiencia: dict | None = None

    # Anti-ICP
    marcadores_anti_icp: list[str] = field(default_factory=list)

    #: {nombre de señal: disponible}. El reporte de lote de esta fila.
    disponibilidad: dict[str, bool] = field(default_factory=dict)
    #: {nombre de señal: por que falta}
    motivos_de_ausencia: dict[str, str] = field(default_factory=dict)

    @property
    def senales_usables(self) -> bool:
        """False si el handle no verifica o el perfil no se leyo."""
        return (
            self.handle_confidence != Confianza.BAJA.value
            and self.estado_perfil == EstadoPerfil.PUBLICO_LEIDO.value
        )

    def a_fila(self) -> dict:
        """Aplanado para la sabana. Las listas y dicts van como JSON."""
        import json

        fila: dict = {}
        for clave, valor in asdict(self).items():
            if isinstance(valor, (list, dict)):
                fila["ig_" + clave] = json.dumps(valor, ensure_ascii=False, default=str) \
                    if valor else None
            else:
                fila["ig_" + clave] = valor
        return fila


# ── Busqueda de handle ────────────────────────────────────────────────────────

def _candidatos_de_ddg(nombre: str, estado: str, page) -> list[str]:
    """Handles candidatos, en orden de aparicion. **Ninguno esta verificado.**

    Se extraen de los `href` de los resultados, no del HTML completo. El bug
    anterior usaba `re.findall` sobre `page.content()`, que devuelve cualquier
    `instagram.com/...` de la pagina: anuncios, pie de pagina, resultados
    ajenos.
    """
    consultas = [
        'site:instagram.com "%s" realtor %s' % (nombre, estado),
        'site:instagram.com "%s" real estate %s' % (nombre, estado),
        'site:instagram.com %s realtor %s' % (nombre, estado),
    ]
    candidatos: list[str] = []

    for consulta in consultas:
        if len(candidatos) >= MAX_CANDIDATOS:
            break
        url = DDG + "?q=" + urllib.parse.quote_plus(consulta) + "&kl=us-en"
        codigo = extraccion._ir(page, url, timeout_ms=20_000)
        if codigo in (403, 429):
            logger.warning("DuckDuckGo limito la busqueda (HTTP {})", codigo)
            time.sleep(random.uniform(30, 70))
            continue
        time.sleep(random.uniform(*PAUSA_BUSQUEDA))

        try:
            enlaces = page.query_selector_all("a[href*='instagram.com']")
        except Exception:  # noqa: BLE001
            continue

        for el in enlaces:
            href = el.get_attribute("href") or ""
            # DuckDuckGo envuelve los resultados; el real va en uddg=
            if "uddg=" in href:
                partes = urllib.parse.parse_qs(urllib.parse.urlparse(href).query)
                href = (partes.get("uddg") or [href])[0]
            m = re.search(
                r"instagram\.com/([A-Za-z0-9_.]{2,30})(?:/|\?|$)", href
            )
            if not m:
                continue
            handle = m.group(1)
            if handle.lower() in _NO_SON_HANDLES:
                continue
            if handle not in candidatos:
                candidatos.append(handle)
            if len(candidatos) >= MAX_CANDIDATOS:
                break

    return candidatos


# ── Construccion de señales ───────────────────────────────────────────────────

def _disponibilidad(crudo: PerfilCrudo, ver: Verificacion) -> tuple[dict, dict]:
    """Que señales se pudieron acreditar y por que las otras no."""
    estado = crudo.diagnostico.estado
    legible = estado.contenido_legible
    confiable = ver.confianza.usable

    motivo_contenido = ""
    if not legible:
        motivo_contenido = estado.motivo()
    elif not confiable:
        motivo_contenido = (
            "el handle no verifico (%s): sus señales no se usan" % ver.razon
        )

    usable = legible and confiable
    disponible = {
        # Estas dos salen del meta y existen incluso en un perfil privado.
        "seguidores": crudo.seguidores is not None,
        "bio": bool(crudo.bio),
        # Estas exigen leer los posts Y que el handle verifique.
        "captions": usable and bool(crudo.posts),
        "idioma_contenido": usable and bool(crudo.captions()),
        "programas": usable and bool(crudo.posts),
        "menciones": usable and bool(crudo.posts),
        "co_marketing": usable and bool(crudo.posts),
        "cadencia": usable and sum(1 for p in crudo.posts if p.fecha) >= 2,
        "engagement": usable and any(
            p.engagement_absoluto() is not None for p in crudo.posts
        ),
        "geotags": usable and any(p.geotag for p in crudo.posts),
        "audiencia": usable and bool(crudo.comentarios),
        "destacadas": bool(crudo.titulos_de_destacadas),
        "link_de_bio": bool(crudo.link_de_bio),
    }

    motivos = {}
    for nombre, hay in disponible.items():
        if hay:
            continue
        if nombre in ("seguidores", "bio", "destacadas", "link_de_bio"):
            motivos[nombre] = motivo_contenido or "no estaba en el perfil"
        else:
            motivos[nombre] = motivo_contenido or (
                "el perfil se leyo pero no habia datos para esta señal"
            )
    return disponible, motivos


def _armar_senales(
    crudo: PerfilCrudo,
    ver: Verificacion,
    candidatos: list[dict],
    *,
    licencia_conocida: str | None = None,
) -> SenalesInstagram:
    s = SenalesInstagram(
        handle=crudo.handle,
        url=crudo.url,
        estado_perfil=crudo.diagnostico.estado.value,
        motivo_estado=crudo.diagnostico.evidencia,
        handle_confidence=ver.confianza.value,
        razon_confianza=ver.razon,
        candidatos_evaluados=candidatos,
        nombre_perfil=crudo.nombre_perfil,
        bio=crudo.bio,
        seguidores=crudo.seguidores,
        n_posts_declarados=crudo.n_posts_declarados,
        categoria_declarada=crudo.categoria_declarada,
        link_de_bio=crudo.link_de_bio,
        destino_link_de_bio=crudo.destino_link_de_bio,
        titulos_de_destacadas=list(crudo.titulos_de_destacadas),
        licencias_en_bio=extraer_licencias_de_bio(crudo.bio),
    )

    disponible, motivos = _disponibilidad(crudo, ver)
    s.disponibilidad = disponible
    s.motivos_de_ausencia = motivos

    if not (crudo.diagnostico.estado.contenido_legible and ver.confianza.usable):
        return s

    # Captions crudos: la fuente, no los flags.
    s.posts_crudos = [
        {
            "url": p.url,
            "fecha": p.fecha.isoformat() if p.fecha else None,
            "caption": p.caption,
            "tipo": p.tipo.value,
            "likes": p.likes,
            "n_comentarios": p.n_comentarios,
            "geotag": p.geotag,
        }
        for p in crudo.posts
    ]

    captions = crudo.captions()
    perfil_idioma: PerfilDeIdioma = perfilar(captions)
    s.idioma = {
        "n_piezas_totales": perfil_idioma.n_piezas_totales,
        "n_piezas_validas": perfil_idioma.n_piezas_validas,
        "n_espanol": perfil_idioma.n_espanol,
        "n_ingles": perfil_idioma.n_ingles,
        "n_mixto": perfil_idioma.n_mixto,
        "n_indeterminado": perfil_idioma.n_indeterminado,
        "ratio_espanol": perfil_idioma.ratio_espanol,
        "idioma_dominante": perfil_idioma.idioma_dominante.value,
        "cumple_evidencia_matriz": perfil_idioma.cumple_evidencia_matriz,
        "nota_de_evidencia": perfil_idioma.nota_de_evidencia,
        "serie_ratio": perfil_idioma.serie_ratio,
    }

    declara, fragmento = declara_acompanamiento(crudo.bio, captions)
    s.evidencia_ancla_idioma = fragmento
    intensidad, grado, regla = perfil_idioma.intensidad_pq14(
        declara_acompanamiento=declara
    )
    s.intensidad_pq14 = intensidad
    s.grado_pq14 = grado or None
    s.regla_pq14 = regla

    s.programas = {
        nombre: {
            "qualifier": datos["qualifier"],
            "n_posts": datos["n_posts"],
            "primer_post": datos["primer_post"].isoformat() if datos["primer_post"] else None,
            "ultimo_post": datos["ultimo_post"].isoformat() if datos["ultimo_post"] else None,
            "ejemplo": datos["ejemplo"],
        }
        for nombre, datos in programas_agregados(crudo.posts).items()
    }

    red: RedDeMenciones = construir_red(crudo.posts)
    s.red_de_menciones = {
        "top": red.top(10),
        "en_co_marketing": red.en_co_marketing,
        "n_posts_con_menciones": red.n_posts_con_menciones,
    }
    s.candidatas_a_lender = red.candidatas_a_lender
    s.n_posts_co_marketing = sum(1 for p in crudo.posts if p.es_co_marketing)
    s.n_posts_con_produccion = sum(1 for p in crudo.posts if p.senal_de_produccion)

    cad: Cadencia = calcular_cadencia(crudo.posts)
    s.cadencia = {
        "n_posts": cad.n_posts,
        "dias_cubiertos": cad.dias_cubiertos,
        "posts_por_semana": cad.posts_por_semana,
        "mediana_dias_entre_posts": cad.mediana_dias_entre_posts,
        "hueco_maximo_dias": cad.hueco_maximo_dias,
        "n_huecos_mayores_a_21d": cad.n_huecos_mayores_a_21d,
        "pico_tras_hueco": cad.pico_tras_hueco,
        "nota": cad.nota,
    }

    eng: EngagementReal = calcular_engagement(crudo.posts, crudo.seguidores)
    s.engagement = {
        "n_posts_con_datos": eng.n_posts_con_datos,
        "n_posts_totales": eng.n_posts_totales,
        "seguidores": eng.seguidores,
        "engagement_medio": eng.engagement_medio,
        "engagement_mediano": eng.engagement_mediano,
        "tasa_media": eng.tasa_media,
        "tasa_mediana": eng.tasa_mediana,
        "suficiente": eng.suficiente,
        "nota": eng.nota,
    }

    s.geotags = geotags_agregados(crudo.posts)

    if crudo.comentarios:
        audiencia: PerfilDeAudiencia = perfilar_audiencia(
            crudo.comentarios,
            n_posts_con_comentarios=crudo.n_posts_con_comentarios,
        )
        # Guarda redundante antes de guardar nada en disco.
        verificar_anonimato(audiencia)
        s.audiencia = {
            "n_comentarios": audiencia.n_comentarios,
            "n_de_terceros": audiencia.n_de_terceros,
            "n_del_agente": audiencia.n_del_agente,
            "tasa_de_respuesta_del_agente": audiencia.tasa_de_respuesta_del_agente,
            "fricciones": audiencia.fricciones,
            "idioma_terceros": audiencia.idioma_terceros,
            "ratio_espanol_terceros": audiencia.ratio_espanol_terceros,
            "redacciones_aplicadas": audiencia.redacciones_aplicadas,
            "nota_de_evidencia": audiencia.nota_de_evidencia,
        }

    marcadores: list[str] = []
    for p in crudo.posts:
        for m in p.marcadores_anti_icp:
            if m not in marcadores:
                marcadores.append(m)
    s.marcadores_anti_icp = marcadores

    return s


# ── Punto de entrada ──────────────────────────────────────────────────────────

def buscar_instagram(
    *,
    nombre: str,
    apellido: str,
    estado: str,
    page,
    licencia_conocida: str | None = None,
    con_comentarios: bool = True,
) -> SenalesInstagram:
    """Busca, verifica y lee el perfil de Instagram de un realtor.

    Evalua hasta `MAX_CANDIDATOS` handles y **se queda con el primero que
    verifique**, no con el primero que aparezca. Si ninguno verifica, devuelve
    el mejor con `handle_confidence = baja` y sus señales sin usar: un handle de
    confianza baja sigue siendo un punto de partida para revision manual.
    """
    nombre_completo = ("%s %s" % (nombre or "", apellido or "")).strip()
    if not nombre_completo:
        s = SenalesInstagram()
        s.motivo_estado = "sin nombre de realtor"
        return s

    candidatos = _candidatos_de_ddg(nombre_completo, estado or "", page)
    if not candidatos:
        s = SenalesInstagram(estado_perfil=EstadoPerfil.SIN_HANDLE.value)
        s.motivo_estado = "la busqueda no devolvio ningun handle candidato"
        s.disponibilidad = {k: False for k in (
            "seguidores", "bio", "captions", "idioma_contenido", "programas",
            "menciones", "co_marketing", "cadencia", "engagement", "geotags",
            "audiencia", "destacadas", "link_de_bio",
        )}
        s.motivos_de_ausencia = {
            k: "no se encontro handle" for k in s.disponibilidad
        }
        return s

    evaluados: list[dict] = []
    mejor: tuple[PerfilCrudo, Verificacion] | None = None

    for handle in candidatos:
        crudo = extraccion.leer_con_backoff(
            page, handle, con_comentarios=False,
        )
        licencias = extraer_licencias_de_bio(crudo.bio)
        coincide_licencia = (
            licencia_conocida
            and any(l.lstrip("0") == licencia_conocida.lstrip("0") for l in licencias)
        )
        ver = verificar(
            nombre_realtor=nombre_completo,
            handle=handle,
            nombre_perfil=crudo.nombre_perfil,
            bio=crudo.bio,
            captions=crudo.captions() or None,
            licencia_en_bio=licencia_conocida if coincide_licencia else None,
        )
        evaluados.append({
            "handle": handle,
            "estado": crudo.diagnostico.estado.value,
            "confianza": ver.confianza.value,
            "razon": ver.razon,
        })
        logger.info(
            "@{} -> {} / confianza {}", handle,
            crudo.diagnostico.estado.value, ver.confianza.value,
        )

        if mejor is None or _mejor_que(ver, mejor[1]):
            mejor = (crudo, ver)
        if ver.confianza is Confianza.ALTA:
            break
        time.sleep(random.uniform(*PAUSA_BUSQUEDA))

    assert mejor is not None
    crudo, ver = mejor

    # Recien ahora, con el handle elegido y verificado, se paga el costo de
    # leer posts y comentarios. Hacerlo por cada candidato multiplicaria por
    # cinco las peticiones a Instagram.
    if ver.confianza.usable and crudo.handle:
        crudo = extraccion.leer_con_backoff(
            page, crudo.handle, con_comentarios=con_comentarios,
        )
        if crudo.link_de_bio:
            crudo.destino_link_de_bio = extraccion.resolver_destino_de_link(
                page, crudo.link_de_bio,
            )
        ver = verificar(
            nombre_realtor=nombre_completo,
            handle=crudo.handle,
            nombre_perfil=crudo.nombre_perfil,
            bio=crudo.bio,
            captions=crudo.captions() or None,
            licencia_en_bio=licencia_conocida if any(
                l.lstrip("0") == (licencia_conocida or "x").lstrip("0")
                for l in extraer_licencias_de_bio(crudo.bio)
            ) else None,
        )

    return _armar_senales(crudo, ver, evaluados, licencia_conocida=licencia_conocida)


def _mejor_que(nueva: Verificacion, actual: Verificacion) -> bool:
    orden = {Confianza.ALTA: 3, Confianza.MEDIA: 2, Confianza.BAJA: 1}
    return orden[nueva.confianza] > orden[actual.confianza]


# ══════════════════════════════════════════════════════════════════════════════
# BLOQUE 1-BIS · CRUDO PRIMERO, PARSER DESPUES
# ══════════════════════════════════════════════════════════════════════════════
#
# Dos pasadas, separadas a proposito:
#
#   pasada 1  correr_lote()      navegador -> ig_raw/{clave}.json
#   pasada 2  parsear_crudos()   ig_raw/*.json -> ig_signals.csv
#
# La segunda no toca la red. Si el parser tiene un bug, se arregla y se
# re-parsea en segundos. Sin el crudo, cada bug del parser cuesta una pasada
# entera contra Instagram: 1.203 perfiles a 30 s son 10 horas, y cada pasada
# gasta la cuota de una cuenta que se puede bloquear.
#
# Interpretar es lo unico que se puede repetir gratis.

import csv  # noqa: E402
import datetime as dt  # noqa: E402
import json  # noqa: E402
import os  # noqa: E402
import statistics  # noqa: E402
import unicodedata  # noqa: E402
from pathlib import Path  # noqa: E402

RAIZ_SCRAPER = Path(__file__).resolve().parent.parent

DIR_CRUDO = RAIZ_SCRAPER / "output" / "ig_raw"
RUTA_CSV = RAIZ_SCRAPER / "output" / "ig_signals.csv"
RUTA_CHECKPOINT = RAIZ_SCRAPER / "output" / "ig_checkpoint.json"
RUTA_MANIFIESTO = RAIZ_SCRAPER / "output" / "ig_manifiesto.json"
RUTA_EXCLUIDOS = RAIZ_SCRAPER / "output" / "ig_objetivo_excluidos.csv"

#: Ritmo. Un perfil cada 20 a 40 segundos, con variacion aleatoria.
PAUSA_ENTRE_PERFILES_LOTE = (20.0, 40.0)

#: Cada tantos perfiles, una pausa larga. No esta en el brief; esta porque un
#: ritmo perfectamente uniforme durante diez horas es, en si mismo, una firma.
CADA_N_PAUSA_LARGA = 40
PAUSA_LARGA = (180.0, 420.0)

#: Separador de piezas en las columnas de texto crudo del CSV.
SEP = " ¶ "

#: Tope por celda. Excel aguanta 32.767; se corta en 30.000 y se declara.
MAX_CHARS_CELDA = 30_000

ESQUEMA_CSV = "ig-signals-v1"


# ── Estados de perfil que expone el CSV ───────────────────────────────────────
#
# El brief pide estos cinco valores exactos. `EstadoPerfil` del modulo tiene
# seis, porque distingue "no habia handle candidato" de "el handle no existe".
# Se mapea, y el detalle fino queda en el JSON crudo.

MAPA_ESTADO_CSV = {
    EstadoPerfil.PUBLICO_LEIDO.value: "publico_leido",
    EstadoPerfil.PRIVADO.value: "privado",
    EstadoPerfil.NO_ENCONTRADO.value: "no_encontrado",
    EstadoPerfil.BLOQUEADO.value: "bloqueado",
    EstadoPerfil.HANDLE_EQUIVOCADO.value: "handle_dudoso",
    EstadoPerfil.SIN_HANDLE.value: "no_encontrado",
}

MAPA_CONFIANZA_CSV = {
    Confianza.ALTA.value: "alta",
    Confianza.MEDIA.value: "media",
    Confianza.BAJA.value: "baja",
}


# ── Lexicos de derivacion ─────────────────────────────────────────────────────
#
# Todos con limite de palabra. `\bestates?\b` matchea "real estate" y no esta
# en ninguna de estas listas: ese patron clasifico como perfil de lujo a mas de
# mil agentes antes de detectarse.

LEX_ITIN = (r"\bitin\b", r"\btax\s?id\b", r"\bsin (?:seguro social|ssn)\b",
            r"\bno ssn\b", r"\bsin numero de seguro\b")
LEX_DPA = (r"\bdpa\b", r"\bdown\s?payment assistance\b", r"\bdown\s?payment help\b",
           r"\bayuda (?:con|para) el enganche\b", r"\bsubvenci[oó]n\b",
           r"\bzero down\b", r"\bsin enganche\b", r"\b0%\s?down\b",
           r"\bgrant\b", r"\benganche\b")
LEX_CREDITO = (r"\bcredit repair\b", r"\bcredit score\b", r"\bbad credit\b",
               r"\breparaci[oó]n de cr[eé]dito\b", r"\bpuntaje de cr[eé]dito\b",
               r"\bsubir tu cr[eé]dito\b", r"\bbankruptcy\b", r"\bforeclosure\b",
               r"\bquiebra\b", r"\bmi cr[eé]dito\b")
LEX_VA = (r"\bva loan\b", r"\bva home\b", r"\bva benefit\b", r"\bpr[eé]stamo va\b",
          r"\bveteran[oa]?s?\b", r"\bmilitar(?:y|es)?\b", r"\bactive duty\b",
          r"\bfunding fee\b", r"\bpcs\b")
LEX_FHA = (r"\bfha\b", r"\bf\.h\.a\.\b", r"\b203\s?k\b")
LEX_PRIMERA_CASA = (r"\bfirst[\s-]?time (?:home\s?)?buyer\b", r"\bprimera casa\b",
                    r"\bprimer hogar\b", r"\bstop renting\b",
                    r"\bdeja(?:r)? de rentar\b", r"\bfirst home\b",
                    r"\bprimera vivienda\b")

#: Pregunta de calificacion en un comentario de tercero. Es la categoria S8.
LEX_PREGUNTA_CALIFICACION = (
    # enganche
    r"\benganch\w*\b", r"\bdown\s?payment\b", r"\bhow much (?:do i need|down)\b",
    r"\bcuanto (?:necesito|hay que|se necesita)\b", r"\bcuanto de entrada\b",
    # credito
    r"\bcr[eé]dito\b", r"\bcredit\b", r"\bscore\b", r"\bpuntaje\b",
    # calificacion
    r"\bcalifi\w*\b", r"\bqualify\b", r"\bqualifi\w*\b", r"\bpre[\s-]?approv\w*\b",
    r"\bprecalifi\w*\b", r"\bcomo (?:empiezo|empezar|aplico)\b",
    r"\bhow do i (?:start|qualify|apply)\b",
    # ITIN y documentos
    r"\bitin\b", r"\bsin (?:papeles|seguro social|ssn)\b", r"\bno ssn\b",
    r"\bdocumento\w*\b", r"\bdocuments?\b", r"\bw2\b", r"\b1099\b",
    r"\btaxes?\b", r"\bimpuestos\b", r"\bpermiso de trabajo\b", r"\bwork permit\b",
)

#: Pistas de que una cuenta etiquetada es hipotecaria o de un loan officer.
#: Es una PISTA, no un hecho: hay que confirmarla contra NMLS antes de
#: escribirla en un dossier. El handle de alguien no prueba su oficio.
PISTAS_HIPOTECARIAS = (
    "loan", "lend", "mortgage", "hipotec", "prestamo", "nmls", "financ",
    "credito", "credit", "escrow", "title", "lo_", "_lo", "loanofficer",
    "homeloan", "mortg", "funding", "capital", "bank",
)

#: Designaciones profesionales. Alimentan S10.
DESIGNACIONES = ("abr", "mrp", "gri", "sres", "crs", "e-pro", "epro", "cips",
                 "nahrep", "ahwd", "rene", "psa", "sfr", "c2ex", "gree",
                 "clhms", "crb")


def _sin_acentos(texto: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFD", texto or "")
        if unicodedata.category(c) != "Mn"
    )


def _cuenta_patrones(textos: list[str], patrones: tuple[str, ...]) -> int:
    """Cuantos textos de la lista disparan al menos un patron."""
    n = 0
    for texto in textos:
        plano = _sin_acentos(texto)
        if any(re.search(p, plano, re.IGNORECASE) for p in patrones):
            n += 1
    return n


# ── La lista de objetivo ──────────────────────────────────────────────────────

RUTA_LIBRO_PACS = (
    RAIZ_SCRAPER.parent / "Data_inputIA" / "Homesi_Scoring_Realtors_v3_PACS (2).xlsx"
)
HOJA_PACS = "Realtors PACS"


@dataclass
class Objetivo:
    """Una fila de la lista a raspar."""

    email: str | None
    nombre: str
    estado: str | None
    handle_del_libro: str | None
    nivel: str | None
    tier: str | None
    licencia: str | None = None

    @property
    def clave(self) -> str:
        """Nombre del archivo crudo: email o, si no hay, handle.

        Se sanea porque va a ser un nombre de archivo en Windows.
        """
        base = (self.email or "").strip().lower()
        if not base:
            base = (self.handle_del_libro or "").strip().lower().lstrip("@")
        if not base:
            base = _sin_acentos(self.nombre).strip().lower().replace(" ", "_")
        return re.sub(r"[^a-z0-9._@-]+", "_", base)[:120] or "sin_clave"


def cargar_objetivo(
    ruta: Path | None = None,
    *,
    excluir_descartados: bool = False,
) -> tuple[list[Objetivo], dict]:
    """Los realtors a raspar, desde el libro de scoring v3 PACS.

    Filtro del brief, literal: `nivel_de_calificacion = MQL` **o** `TIER` que
    empiece por A o por B.

    Devuelve (lista, informe). El informe trae la composicion, porque el filtro
    literal tiene una consecuencia que conviene ver antes de lanzar diez horas
    de scraping. Ver el aviso de abajo.
    """
    import openpyxl

    ruta = ruta or RUTA_LIBRO_PACS
    if not ruta.exists():
        raise SystemExit(
            "No encuentro el libro de scoring en:\n  %s\n\n"
            "Es el insumo de la lista de objetivo. Vive en Data_inputIA/, que "
            "esta fuera del control de versiones." % ruta
        )

    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    if HOJA_PACS not in wb.sheetnames:
        wb.close()
        raise SystemExit(
            "El libro no tiene la hoja %r. Hojas: %s" % (HOJA_PACS, wb.sheetnames)
        )
    ws = wb[HOJA_PACS]
    filas = list(ws.iter_rows(values_only=True))
    wb.close()

    enc = [("" if c is None else str(c)).strip() for c in filas[0]]

    def columna(*fragmentos: str) -> str | None:
        for c in enc:
            bajo = c.lower()
            if all(f.lower() in bajo for f in fragmentos):
                return c
        return None

    c_nivel = columna("nivel_de_calificacion")
    c_tier = columna("tier")
    c_email = columna("email")
    c_handle = columna("ig", "handle")
    c_nombre = columna("nombre")
    c_estado = columna("estado")
    if not (c_nivel and c_tier and c_nombre):
        raise SystemExit(
            "Faltan columnas clave en %r. nivel=%r tier=%r nombre=%r"
            % (HOJA_PACS, c_nivel, c_tier, c_nombre)
        )

    def texto(fila: dict, clave: str | None) -> str:
        if not clave:
            return ""
        v = fila.get(clave)
        return ("" if v is None else str(v)).strip()

    registros = []
    for fila in filas[1:]:
        if all(c is None or str(c).strip() == "" for c in fila):
            continue
        registros.append({
            enc[i]: fila[i]
            for i in range(min(len(enc), len(fila)))
            if enc[i]
        })

    objetivos: list[Objetivo] = []
    excluidos: list[dict] = []
    composicion_nivel: dict[str, int] = {}
    composicion_tier: dict[str, int] = {}

    for r in registros:
        nivel = texto(r, c_nivel).upper()
        tier = texto(r, c_tier).upper()
        if not (nivel == "MQL" or tier.startswith("A") or tier.startswith("B")):
            continue

        composicion_nivel[nivel or "(vacio)"] = (
            composicion_nivel.get(nivel or "(vacio)", 0) + 1
        )
        composicion_tier[tier or "(vacio)"] = (
            composicion_tier.get(tier or "(vacio)", 0) + 1
        )

        obj = Objetivo(
            email=texto(r, c_email) or None,
            nombre=texto(r, c_nombre),
            estado=texto(r, c_estado) or None,
            handle_del_libro=texto(r, c_handle) or None,
            nivel=nivel or None,
            tier=tier or None,
        )

        # El filtro literal re-admite gente que la metodologia ya excluyo: hay
        # registros DESCARTADO y RECLASIFICADO cuyo TIER es A o B. DESCARTADO
        # significa "no contactar" -- perfil de lujo, inversion pura, o por
        # debajo del piso de volumen -- y RECLASIFICADO POR MMI significa que
        # los datos transaccionales lo sacaron del ICP.
        #
        # Raspar sus perfiles no hace daño, pero meterlos en un CSV que
        # alimenta prospeccion si. Se separan a un archivo aparte y se declaran.
        es_excluido_por_metodologia = (
            nivel.startswith("DESCARTADO") or nivel.startswith("RECLASIFICADO")
        )
        if es_excluido_por_metodologia:
            excluidos.append({
                "email": obj.email, "nombre": obj.nombre, "estado": obj.estado,
                "handle": obj.handle_del_libro, "nivel": nivel, "tier": tier,
                "motivo": ("entra por TIER %s pero la metodologia ya lo excluyo "
                           "con nivel %s" % (tier, nivel)),
            })
            if excluir_descartados:
                continue

        objetivos.append(obj)

    informe = {
        "libro": str(ruta),
        "hoja": HOJA_PACS,
        "filas_en_la_hoja": len(registros),
        "objetivo": len(objetivos),
        "con_handle_en_el_libro": sum(1 for o in objetivos if o.handle_del_libro),
        "sin_handle_en_el_libro": sum(1 for o in objetivos if not o.handle_del_libro),
        "con_email": sum(1 for o in objetivos if o.email),
        "composicion_por_nivel": dict(sorted(composicion_nivel.items())),
        "composicion_por_tier": dict(sorted(composicion_tier.items())),
        "excluidos_por_metodologia": excluidos,
        "excluir_descartados_aplicado": excluir_descartados,
        "horas_estimadas_a_30s": round(len(objetivos) * 30 / 3600, 1),
    }
    return objetivos, informe


# ── Persistencia del crudo ────────────────────────────────────────────────────

def guardar_crudo(objetivo: Objetivo, crudo: dict, *, dir_crudo: Path | None = None
                  ) -> Path:
    """Escribe ig_raw/{clave}.json. Un archivo por perfil, con timestamp dentro.

    Se escribe a un temporal y se renombra, para que una interrupcion a mitad
    de escritura no deje un JSON truncado que el parser lea como valido.
    """
    dir_crudo = dir_crudo or DIR_CRUDO
    dir_crudo.mkdir(parents=True, exist_ok=True)

    crudo = dict(crudo)
    crudo["objetivo"] = {
        "email": objetivo.email,
        "nombre": objetivo.nombre,
        "estado": objetivo.estado,
        "handle_del_libro": objetivo.handle_del_libro,
        "nivel": objetivo.nivel,
        "tier": objetivo.tier,
    }

    destino = dir_crudo / ("%s.json" % objetivo.clave)
    temporal = destino.with_suffix(".json.tmp")
    with temporal.open("w", encoding="utf-8") as fh:
        json.dump(crudo, fh, ensure_ascii=False, indent=1, default=str)
    os.replace(temporal, destino)
    return destino


def leer_crudo(ruta: Path) -> dict | None:
    try:
        with ruta.open("r", encoding="utf-8") as fh:
            return json.load(fh)
    except (json.JSONDecodeError, OSError) as exc:
        logger.warning("crudo ilegible {}: {}", ruta.name, exc)
        return None


# ── Checkpoint ────────────────────────────────────────────────────────────────

def _cargar_checkpoint(ruta: Path | None = None) -> dict:
    ruta = ruta or RUTA_CHECKPOINT
    if not ruta.exists():
        return {"hechos": [], "detenido_por": None, "iniciado_en": None}
    try:
        with ruta.open("r", encoding="utf-8") as fh:
            return json.load(fh)
    except (json.JSONDecodeError, OSError):
        return {"hechos": [], "detenido_por": None, "iniciado_en": None}


def _guardar_checkpoint(estado: dict, ruta: Path | None = None) -> None:
    ruta = ruta or RUTA_CHECKPOINT
    ruta.parent.mkdir(parents=True, exist_ok=True)
    temporal = ruta.with_suffix(".json.tmp")
    with temporal.open("w", encoding="utf-8") as fh:
        json.dump(estado, fh, ensure_ascii=False, indent=1)
    os.replace(temporal, ruta)


# ── Pasada 1 · el lote ────────────────────────────────────────────────────────

#: Estados que hacen que el lote se DETENGA, no que reintente.
#:
#: El brief lo pide explicito: "Si aparece un desafio, un captcha o un bloqueo:
#: detente y registra. No reintentes en bucle." Reintentar ante un bloqueo es
#: como se pasa de un desafio temporal a una cuenta suspendida.
ESTADOS_QUE_DETIENEN = {EstadoPerfil.BLOQUEADO.value}

#: Bloqueos consecutivos tolerados antes de cortar. Uno puede ser un perfil
#: raro; tres seguidos es la cuenta.
MAX_BLOQUEOS_SEGUIDOS = 3


def correr_lote(
    *,
    limite: int | None = None,
    reanudar: bool = True,
    excluir_descartados: bool = False,
    headless: bool = True,
    dir_crudo: Path | None = None,
    con_comentarios: bool = True,
    pausa_perfiles: tuple[float, float] = PAUSA_ENTRE_PERFILES_LOTE,
) -> dict:
    """Pasada 1: raspa y guarda el crudo. No deriva nada.

    Reanudable: cada perfil se marca en el checkpoint en cuanto su JSON esta en
    disco, asi que una interrupcion cuesta un perfil y no el lote.
    """
    try:
        try:
            from patchright.sync_api import sync_playwright
        except ImportError:
            from playwright.sync_api import sync_playwright
    except ImportError as exc:
        raise SystemExit(
            "Falta el navegador. La pasada 1 lo necesita; la pasada 2 (--parsear)\n"
            "no, y es la que se puede correr en cualquier maquina.\n"
            "\n"
            "    pip install -r realtor_scraper/requirements.txt\n"
            "    python -m playwright install chromium\n"
            "\n"
            "Detalle: %s" % exc
        ) from exc

    from navegador import (
        PERFIL_IG_DIR,
        crear_contexto,
        sesion_de_instagram_iniciada,
    )

    dir_crudo = dir_crudo or DIR_CRUDO
    objetivos, informe = cargar_objetivo(excluir_descartados=excluir_descartados)

    checkpoint = _cargar_checkpoint() if reanudar else {
        "hechos": [], "detenido_por": None, "iniciado_en": None
    }
    hechos = set(checkpoint.get("hechos") or [])
    if not checkpoint.get("iniciado_en"):
        checkpoint["iniciado_en"] = dt.datetime.now().isoformat(timespec="seconds")

    pendientes = [o for o in objetivos if o.clave not in hechos]
    if limite:
        pendientes = pendientes[:limite]

    logger.info("Objetivo: {} · ya hechos: {} · pendientes en esta corrida: {}",
                len(objetivos), len(hechos), len(pendientes))
    logger.info("Ritmo: {:.0f}-{:.0f} s por perfil -> ~{:.1f} h para los pendientes",
                pausa_perfiles[0], pausa_perfiles[1],
                len(pendientes) * sum(pausa_perfiles) / 2 / 3600)

    resumen = {
        "esquema": ESQUEMA_CSV,
        "iniciado_en": dt.datetime.now().isoformat(timespec="seconds"),
        "objetivo_total": len(objetivos),
        "pendientes_al_inicio": len(pendientes),
        "procesados": 0,
        "por_estado": {},
        "detenido_por": None,
        "informe_objetivo": informe,
        "motor_idioma": None,
    }

    if not pendientes:
        logger.info("No hay pendientes. Nada que hacer.")
        resumen["terminado_en"] = dt.datetime.now().isoformat(timespec="seconds")
        return resumen

    bloqueos_seguidos = 0

    with sync_playwright() as p:
        contexto = crear_contexto(p, headless=headless, perfil_dir=PERFIL_IG_DIR)
        page = contexto.new_page()

        hay_sesion, evidencia = sesion_de_instagram_iniciada(page)
        resumen["sesion"] = {"iniciada": hay_sesion, "evidencia": evidencia}
        if not hay_sesion:
            contexto.close()
            resumen["detenido_por"] = "sin_sesion"
            resumen["terminado_en"] = dt.datetime.now().isoformat(timespec="seconds")
            logger.error(
                "NO hay sesion de Instagram ({}). Sin sesion se lee una version "
                "recortada del perfil: bio y contadores si, posts y comentarios "
                "no. Eso no da error, da un lote entero de perfiles que parecen "
                "vacios.\n\n"
                "    python -m instagram.finder --iniciar-sesion\n", evidencia,
            )
            return resumen

        logger.info("Sesion OK: {}", evidencia)

        for i, objetivo in enumerate(pendientes, start=1):
            handle = (objetivo.handle_del_libro or "").strip().lstrip("@")
            if not handle:
                # Sin handle en el libro se busca, que es lo caro. Son 3 de 1.203.
                senales = buscar_instagram(
                    nombre=objetivo.nombre.split(" ")[0] if objetivo.nombre else "",
                    apellido=" ".join(objetivo.nombre.split(" ")[1:])
                    if objetivo.nombre else "",
                    estado=objetivo.estado or "",
                    page=page,
                    con_comentarios=False,
                )
                handle = senales.handle or ""

            if handle:
                crudo = extraccion.capturar_crudo(
                    page, handle, con_comentarios=con_comentarios,
                )
            else:
                crudo = {
                    "esquema": "ig-crudo-v1",
                    "handle_pedido": None,
                    "capturado_en": dt.datetime.now(dt.timezone.utc).isoformat(
                        timespec="seconds"
                    ),
                    "estado_perfil": EstadoPerfil.SIN_HANDLE.value,
                    "estado_evidencia": "ni el libro ni la busqueda dieron handle",
                    "perfil": None, "posts": [], "comentarios_por_post": {},
                    "errores": [],
                }

            guardar_crudo(objetivo, crudo, dir_crudo=dir_crudo)
            hechos.add(objetivo.clave)
            checkpoint["hechos"] = sorted(hechos)
            _guardar_checkpoint(checkpoint)

            estado = crudo.get("estado_perfil") or "desconocido"
            resumen["por_estado"][estado] = resumen["por_estado"].get(estado, 0) + 1
            resumen["procesados"] += 1

            logger.info("[{}/{}] {} -> @{} · {}",
                        i, len(pendientes), objetivo.nombre[:28],
                        handle or "-", estado)

            if estado in ESTADOS_QUE_DETIENEN:
                bloqueos_seguidos += 1
                if bloqueos_seguidos >= MAX_BLOQUEOS_SEGUIDOS:
                    checkpoint["detenido_por"] = "bloqueo"
                    checkpoint["detenido_en"] = dt.datetime.now().isoformat(
                        timespec="seconds"
                    )
                    _guardar_checkpoint(checkpoint)
                    resumen["detenido_por"] = "bloqueo"
                    logger.error(
                        "{} bloqueos seguidos. ME DETENGO y no reintento.\n"
                        "Ultima evidencia: {}\n"
                        "El checkpoint quedo en {}: al reanudar sigue desde aca.\n"
                        "Antes de reanudar, esperar varias horas y confirmar a "
                        "mano que la cuenta no tiene un desafio pendiente.",
                        bloqueos_seguidos, crudo.get("estado_evidencia"),
                        RUTA_CHECKPOINT,
                    )
                    break
            else:
                bloqueos_seguidos = 0

            if i < len(pendientes):
                if i % CADA_N_PAUSA_LARGA == 0:
                    larga = random.uniform(*PAUSA_LARGA)
                    logger.info("Pausa larga de {:.0f} s tras {} perfiles", larga, i)
                    time.sleep(larga)
                else:
                    time.sleep(random.uniform(*pausa_perfiles))

        contexto.close()

    resumen["terminado_en"] = dt.datetime.now().isoformat(timespec="seconds")
    RUTA_MANIFIESTO.parent.mkdir(parents=True, exist_ok=True)
    with RUTA_MANIFIESTO.open("w", encoding="utf-8") as fh:
        json.dump(resumen, fh, ensure_ascii=False, indent=2, default=str)
    return resumen


# ── Pasada 2 · el parser ──────────────────────────────────────────────────────

def _fecha_de_post(post: dict) -> dt.date | None:
    ts = post.get("timestamp")
    if ts:
        try:
            return dt.datetime.fromtimestamp(int(ts), dt.timezone.utc).date()
        except (ValueError, OSError, OverflowError):
            pass
    iso = post.get("fecha_iso")
    if iso:
        try:
            return dt.date.fromisoformat(str(iso)[:10])
        except ValueError:
            pass
    return None


def _tipo_de_post(post: dict) -> str:
    typename = (post.get("typename") or "").lower()
    if "video" in typename or post.get("es_video"):
        # Un GraphVideo con duracion corta y sin ser sidecar es un reel en la
        # practica. Instagram no distingue reel de video en este campo.
        return "reel"
    if "sidecar" in typename or (post.get("n_hijos") or 0) > 1:
        return "carrusel"
    if "image" in typename:
        return "foto"
    return typename or "desconocido"


def _menciones_de_caption(caption: str | None) -> list[str]:
    if not caption:
        return []
    return list(dict.fromkeys(
        m.lower() for m in re.findall(r"@([A-Za-z0-9_.]{2,30})", caption)
    ))


def _es_hipotecaria(handle: str) -> bool:
    bajo = (handle or "").lower()
    return any(p in bajo for p in PISTAS_HIPOTECARIAS)


def _designaciones_de_bio(bio: str | None) -> list[str]:
    if not bio:
        return []
    plano = _sin_acentos(bio).lower()
    salida = []
    for d in DESIGNACIONES:
        if re.search(r"\b" + re.escape(d) + r"\b", plano):
            etiqueta = d.upper().replace("EPRO", "e-PRO")
            if etiqueta not in salida:
                salida.append(etiqueta)
    return salida


def _armar_texto_captions(posts_ordenados: list[tuple[dt.date | None, str]]) -> str:
    """`2026-08-14 | caption ¶ 2026-08-02 | otro`, cronologico inverso."""
    piezas = []
    for fecha, caption in posts_ordenados:
        limpio = re.sub(r"\s+", " ", (caption or "")).strip()
        if not limpio:
            continue
        prefijo = fecha.isoformat() if fecha else "sin-fecha"
        piezas.append("%s | %s" % (prefijo, limpio))
    return SEP.join(piezas)


def _armar_texto_comentarios(textos: list[str]) -> str:
    """Solo el texto, sin el handle del autor. Separados por pilcrow.

    El handle del comentarista vive en el JSON crudo si alguna vez hace falta.
    En el CSV no, porque son terceros que no son nuestro prospecto.
    """
    piezas = [re.sub(r"\s+", " ", (t or "")).strip() for t in textos]
    return SEP.join(p for p in piezas if p)


def _truncar(texto: str) -> tuple[str, bool]:
    if len(texto) <= MAX_CHARS_CELDA:
        return texto, False
    return texto[:MAX_CHARS_CELDA], True


def parsear_crudo(crudo: dict) -> dict:
    """Un JSON crudo -> una fila derivada de ig_signals.csv.

    No toca la red. Es la unica funcion que interpreta, y por eso es la unica
    que se puede corregir y volver a correr gratis.
    """
    from instagram.idioma import (
        MIN_LETRAS_COMENTARIO,
        Idioma,
        clasificar_comentario,
        clasificar_pieza,
        motor_en_uso,
    )

    objetivo = crudo.get("objetivo") or {}
    perfil = crudo.get("perfil") or {}
    posts = crudo.get("posts") or []
    comentarios_por_post = crudo.get("comentarios_por_post") or {}

    estado_bruto = crudo.get("estado_perfil") or EstadoPerfil.SIN_HANDLE.value
    estado_csv = MAPA_ESTADO_CSV.get(estado_bruto, estado_bruto)

    handle = perfil.get("handle") or crudo.get("handle_pedido") or \
        objetivo.get("handle_del_libro")
    bio = perfil.get("bio")
    seguidores = perfil.get("seguidores")

    # ── Verificacion del handle ───────────────────────────────────────────────
    captions_todos = [p.get("caption") for p in posts if p.get("caption")]
    ver = verificar(
        nombre_realtor=objetivo.get("nombre") or "",
        handle=handle,
        nombre_perfil=perfil.get("nombre_visible"),
        bio=bio,
        captions=captions_todos or None,
    )
    confianza_csv = MAPA_CONFIANZA_CSV.get(ver.confianza.value, "baja")

    # Si el handle no verifica, el estado pasa a handle_dudoso y sus señales de
    # contenido NO se derivan. Un handle equivocado no produce un error:
    # produce catorce señales sobre la persona equivocada.
    senales_usables = (
        estado_bruto == EstadoPerfil.PUBLICO_LEIDO.value
        and ver.confianza is not Confianza.BAJA
    )
    if (estado_bruto == EstadoPerfil.PUBLICO_LEIDO.value
            and ver.confianza is Confianza.BAJA):
        estado_csv = "handle_dudoso"

    fila: dict = {
        "email": objetivo.get("email"),
        "nombre": objetivo.get("nombre"),
        "estado": objetivo.get("estado"),
        "handle": handle,
        "estado_perfil": estado_csv,
        "handle_confianza": confianza_csv,
    }

    #: Todas las señales de contenido arrancan en None. Nunca en 0 ni en False.
    #: Un perfil privado no es un perfil sin español: es un perfil que no leimos.
    for col in ("captions_n", "captions_es_ratio", "comentarios_n",
                "comentarios_es_ratio", "menciona_lender", "menciona_itin",
                "menciona_dpa", "menciona_credito", "menciona_va",
                "menciona_fha", "menciona_primera_casa",
                "comentarios_pregunta_calificacion",
                "cuentas_hipotecarias_etiquetadas", "posts_comarketing",
                "tipo_post_reel_pct", "dias_entre_posts_mediana",
                "hueco_max_dias", "engagement_rate", "geotags_top",
                "captions_texto", "comentarios_texto", "comentarios_del_agente"):
        fila[col] = None

    # Estas dos se leen del perfil y existen incluso en una cuenta privada.
    fila["designaciones"] = ", ".join(_designaciones_de_bio(bio)) or None
    fila["destacadas_titulos"] = ", ".join(
        perfil.get("titulos_destacadas") or []
    ) or None
    fila["texto_truncado"] = False

    if not senales_usables:
        return fila

    # ── Captions ──────────────────────────────────────────────────────────────
    con_fecha = [(_fecha_de_post(p), p.get("caption") or "") for p in posts]
    # Cronologico inverso para el texto; cronologico para la serie de idioma.
    inverso = sorted(
        con_fecha,
        key=lambda par: (par[0] or dt.date.min),
        reverse=True,
    )

    captions_n = sum(1 for _, c in con_fecha if (c or "").strip())
    fila["captions_n"] = captions_n

    clasificadas = [clasificar_pieza(c) for _, c in con_fecha if (c or "").strip()]
    n_es = sum(1 for x in clasificadas if x.idioma is Idioma.ESPANOL)
    # El brief define el denominador: "sobre captions_n". Se respeta.
    # El desglose de clasificables queda en el JSON crudo y en el manifiesto.
    fila["captions_es_ratio"] = (
        round(n_es / captions_n, 4) if captions_n else None
    )

    # ── Comentarios ───────────────────────────────────────────────────────────
    handle_agente = (handle or "").strip().lstrip("@").lower()
    terceros: list[str] = []
    del_agente: list[str] = []
    for comentarios in comentarios_por_post.values():
        for c in comentarios or []:
            texto = (c.get("texto") or "").strip()
            if not texto:
                continue
            autor = (c.get("autor_handle") or "").strip().lstrip("@").lower()
            if autor and autor == handle_agente:
                del_agente.append(texto)
            else:
                terceros.append(texto)

    fila["comentarios_n"] = len(terceros)
    if terceros:
        clas_com = [clasificar_comentario(t) for t in terceros]
        n_es_com = sum(1 for x in clas_com if x.idioma is Idioma.ESPANOL)
        fila["comentarios_es_ratio"] = round(n_es_com / len(terceros), 4)
    else:
        fila["comentarios_es_ratio"] = None

    fila["comentarios_pregunta_calificacion"] = _cuenta_patrones(
        terceros, LEX_PREGUNTA_CALIFICACION
    )

    # ── Programas mencionados · conteo de posts que lo mencionan ─────────────
    textos_contenido = [c for _, c in con_fecha if (c or "").strip()]
    if bio:
        textos_contenido_con_bio = textos_contenido + [bio]
    else:
        textos_contenido_con_bio = textos_contenido

    fila["menciona_itin"] = _cuenta_patrones(textos_contenido_con_bio, LEX_ITIN)
    fila["menciona_dpa"] = _cuenta_patrones(textos_contenido_con_bio, LEX_DPA)
    fila["menciona_credito"] = _cuenta_patrones(textos_contenido_con_bio, LEX_CREDITO)
    fila["menciona_va"] = _cuenta_patrones(textos_contenido_con_bio, LEX_VA)
    fila["menciona_fha"] = _cuenta_patrones(textos_contenido_con_bio, LEX_FHA)
    fila["menciona_primera_casa"] = _cuenta_patrones(
        textos_contenido_con_bio, LEX_PRIMERA_CASA
    )

    # ── S6 · cuentas hipotecarias etiquetadas ────────────────────────────────
    #
    # Es el campo mas valioso del archivo. S6 -- con quien trabaja -- esta vacia
    # en las 4.249 filas, y si el agente etiqueta a un loan officer o a una
    # hipotecaria en sus posts, esta diciendo con quien trabaja.
    etiquetadas: dict[str, int] = {}
    for p in posts:
        for cuenta in (p.get("etiquetadas") or []):
            k = str(cuenta).lower().lstrip("@")
            etiquetadas[k] = etiquetadas.get(k, 0) + 1
        for cuenta in _menciones_de_caption(p.get("caption")):
            etiquetadas[cuenta] = etiquetadas.get(cuenta, 0) + 1

    hipotecarias = sorted(
        ((h, n) for h, n in etiquetadas.items() if _es_hipotecaria(h)),
        key=lambda kv: (-kv[1], kv[0]),
    )
    fila["cuentas_hipotecarias_etiquetadas"] = ", ".join(
        "@%s (%d)" % (h, n) for h, n in hipotecarias
    ) or None
    # menciona_lender: el NOMBRE de la cuenta, no un booleano. Es lo que pide el
    # brief, y el nombre es lo unico que sirve para abrir la conversacion.
    fila["menciona_lender"] = ("@%s" % hipotecarias[0][0]) if hipotecarias else None

    # ── Co-marketing ─────────────────────────────────────────────────────────
    lex_comarketing = (
        r"\bco[\s-]?host(?:ed|ing)?\b", r"\bworkshop\b", r"\btaller\b",
        r"\bseminar(?:io)?\b", r"\bwebinar\b", r"\bcharla\b",
        r"\bclase gratis\b", r"\bfree class\b", r"\bevento\b",
        r"\bpartnered with\b", r"\ben conjunto con\b", r"\bjunto a\b",
        r"\bcolabora\w* con\b", r"\bhomebuyer (?:class|workshop|seminar)\b",
    )
    n_comarketing = 0
    for p in posts:
        caption = p.get("caption") or ""
        tiene_socio = bool(_menciones_de_caption(caption)
                           or (p.get("etiquetadas") or []))
        if tiene_socio and _cuenta_patrones([caption], lex_comarketing):
            n_comarketing += 1
    fila["posts_comarketing"] = n_comarketing

    # ── Cadencia y tipo ──────────────────────────────────────────────────────
    tipos = [_tipo_de_post(p) for p in posts]
    fila["tipo_post_reel_pct"] = (
        round(sum(1 for t in tipos if t == "reel") / len(tipos) * 100, 1)
        if tipos else None
    )

    fechas = sorted(f for f, _ in con_fecha if f)
    if len(fechas) >= 2:
        deltas = [(fechas[i + 1] - fechas[i]).days for i in range(len(fechas) - 1)]
        fila["dias_entre_posts_mediana"] = float(statistics.median(deltas))
        fila["hueco_max_dias"] = max(deltas)
    else:
        fila["dias_entre_posts_mediana"] = None
        fila["hueco_max_dias"] = None

    # ── Engagement real ──────────────────────────────────────────────────────
    #
    # (likes + comentarios) / seguidores. El proxy anterior era seguidores, que
    # correlaciona +0,0014 con unidades vendidas: nada.
    pares = [
        (p.get("likes"), p.get("n_comentarios")) for p in posts
        if p.get("likes") is not None and p.get("n_comentarios") is not None
    ]
    if pares and seguidores:
        medio = statistics.fmean(l + c for l, c in pares)
        fila["engagement_rate"] = round(medio / seguidores, 5)
    else:
        fila["engagement_rate"] = None

    # ── Geotags ──────────────────────────────────────────────────────────────
    geos: dict[str, int] = {}
    for p in posts:
        ubic = p.get("ubicacion") or {}
        nombre_geo = (ubic.get("nombre") or "").strip()
        if nombre_geo:
            geos[nombre_geo] = geos.get(nombre_geo, 0) + 1
    fila["geotags_top"] = ", ".join(
        "%s (%d)" % (g, n)
        for g, n in sorted(geos.items(), key=lambda kv: (-kv[1], kv[0]))[:5]
    ) or None

    # ── Texto crudo en el CSV ────────────────────────────────────────────────
    t_captions, trunc1 = _truncar(_armar_texto_captions(inverso))
    t_terceros, trunc2 = _truncar(_armar_texto_comentarios(terceros))
    t_agente, trunc3 = _truncar(_armar_texto_comentarios(del_agente))
    fila["captions_texto"] = t_captions or None
    fila["comentarios_texto"] = t_terceros or None
    fila["comentarios_del_agente"] = t_agente or None
    fila["texto_truncado"] = bool(trunc1 or trunc2 or trunc3)

    return fila


#: Las columnas del CSV, en orden. **Contrato con la app que lo consume.**
#: No se reordenan ni se renombran sin avisar.
COLUMNAS_CSV = [
    "email", "nombre", "estado", "handle", "estado_perfil", "handle_confianza",
    "captions_n", "captions_es_ratio", "comentarios_n", "comentarios_es_ratio",
    "menciona_lender", "menciona_itin", "menciona_dpa", "menciona_credito",
    "menciona_va", "menciona_fha", "menciona_primera_casa",
    "comentarios_pregunta_calificacion",
    "cuentas_hipotecarias_etiquetadas", "posts_comarketing",
    "tipo_post_reel_pct", "dias_entre_posts_mediana", "hueco_max_dias",
    "engagement_rate", "designaciones", "destacadas_titulos", "geotags_top",
    # Texto crudo, agregado en el Bloque 1-bis.
    "captions_texto", "comentarios_texto", "comentarios_del_agente",
    "texto_truncado",
]


def parsear_crudos(
    *,
    dir_crudo: Path | None = None,
    ruta_csv: Path | None = None,
) -> dict:
    """Pasada 2: ig_raw/*.json -> ig_signals.csv. Sin red.

    UTF-8 **con BOM** para que Excel no destroce los acentos, y comillas dobles
    escapadas segun CSV estandar.
    """
    from instagram.idioma import motor_en_uso

    dir_crudo = dir_crudo or DIR_CRUDO
    ruta_csv = ruta_csv or RUTA_CSV

    if not dir_crudo.exists():
        raise SystemExit(
            "No hay crudos en %s. Corre primero la pasada 1:\n"
            "    python -m instagram.finder --lote" % dir_crudo
        )

    archivos = sorted(dir_crudo.glob("*.json"))
    if not archivos:
        raise SystemExit("No hay archivos .json en %s." % dir_crudo)

    filas = []
    ilegibles = []
    for ruta in archivos:
        crudo = leer_crudo(ruta)
        if crudo is None:
            ilegibles.append(ruta.name)
            continue
        try:
            filas.append(parsear_crudo(crudo))
        except Exception as exc:  # noqa: BLE001
            # Un crudo que rompe el parser NO tumba la pasada: se registra y se
            # sigue. El crudo esta en disco, asi que se re-parsea gratis.
            ilegibles.append("%s (parser: %r)" % (ruta.name, exc))

    ruta_csv.parent.mkdir(parents=True, exist_ok=True)
    with ruta_csv.open("w", encoding="utf-8-sig", newline="") as fh:
        escritor = csv.DictWriter(
            fh, fieldnames=COLUMNAS_CSV, extrasaction="ignore",
            quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n",
        )
        escritor.writeheader()
        for f in filas:
            escritor.writerow({k: f.get(k) for k in COLUMNAS_CSV})

    por_estado: dict[str, int] = {}
    por_confianza: dict[str, int] = {}
    for f in filas:
        e = f.get("estado_perfil") or "?"
        c = f.get("handle_confianza") or "?"
        por_estado[e] = por_estado.get(e, 0) + 1
        por_confianza[c] = por_confianza.get(c, 0) + 1

    con_ratio = [f["captions_es_ratio"] for f in filas
                 if f.get("captions_es_ratio") is not None]
    con_lender = [f for f in filas if f.get("menciona_lender")]
    con_preguntas = [f for f in filas
                     if (f.get("comentarios_pregunta_calificacion") or 0) > 0]

    informe = {
        "esquema": ESQUEMA_CSV,
        "parseado_en": dt.datetime.now().isoformat(timespec="seconds"),
        "motor_idioma": motor_en_uso(),
        "crudos_leidos": len(archivos),
        "filas_escritas": len(filas),
        "ilegibles": ilegibles,
        "csv": str(ruta_csv),
        "por_estado_perfil": dict(sorted(por_estado.items())),
        "por_handle_confianza": dict(sorted(por_confianza.items())),
        "captions_es_ratio": {
            "filas_con_ratio": len(con_ratio),
            "media": round(statistics.fmean(con_ratio), 4) if con_ratio else None,
            "mediana": round(statistics.median(con_ratio), 4) if con_ratio else None,
        },
        "S6_menciona_lender": {
            "filas_con_lender": len(con_lender),
            "de": len(filas),
        },
        "S8_comentarios_pregunta_calificacion": {
            "filas_con_al_menos_una": len(con_preguntas),
            "de": len(filas),
        },
        "truncados": sum(1 for f in filas if f.get("texto_truncado")),
    }
    return informe


# ── La revision de 20 perfiles, antes de lanzar sobre los 1.000 ──────────────

def revisar_piloto(*, dir_crudo: Path | None = None, n: int = 20) -> dict:
    """Imprime lo que hay que mirar A MANO despues del piloto.

    El brief pide tres preguntas concretas y esta funcion las pone adelante en
    vez de dejarlas a criterio de quien mire:

      1. ¿los captions son los reales?
      2. ¿el ratio de español tiene sentido contra lo que se ve en pantalla?
      3. ¿aparecen perfiles privados marcados como privados?

    Y el criterio de parada: si `captions_es_ratio` medio sigue cerca del 3%,
    algo quedo leyendo el alt-text y **hay que detenerse**.
    """
    dir_crudo = dir_crudo or DIR_CRUDO
    archivos = sorted(dir_crudo.glob("*.json"))[:n]
    if not archivos:
        raise SystemExit("No hay crudos en %s." % dir_crudo)

    filas = []
    for ruta in archivos:
        crudo = leer_crudo(ruta)
        if crudo:
            filas.append((ruta.name, crudo, parsear_crudo(crudo)))

    print("=" * 78)
    print("PILOTO · %d perfiles · revision a mano" % len(filas))
    print("=" * 78)

    estados: dict[str, int] = {}
    ratios = []
    for nombre, crudo, fila in filas:
        e = fila.get("estado_perfil") or "?"
        estados[e] = estados.get(e, 0) + 1
        if fila.get("captions_es_ratio") is not None:
            ratios.append(fila["captions_es_ratio"])

    print("")
    print("1 · ESTADOS DE PERFIL")
    for e, n_e in sorted(estados.items()):
        print("   %-16s %d" % (e, n_e))
    if "privado" not in estados:
        print("")
        print("   ⚠ NINGUN perfil privado en %d. Es posible, pero es la misma" % len(filas))
        print("     forma del bug viejo: ig_is_private salia 0 en las 5.620 filas.")
        print("     Confirma a mano abriendo dos o tres perfiles.")

    print("")
    print("2 · CAPTIONS REALES · los tres primeros de cada perfil")
    for nombre, crudo, fila in filas[:6]:
        print("")
        print("   --- %s · @%s · %s ---"
              % (nombre, fila.get("handle"), fila.get("estado_perfil")))
        posts = (crudo.get("posts") or [])[:3]
        if not posts:
            print("      (sin posts: %s)" % (crudo.get("estado_evidencia") or "?"))
        for p in posts:
            caption = re.sub(r"\s+", " ", p.get("caption") or "").strip()
            alt = re.sub(r"\s+", " ", p.get("accesibilidad_alt") or "").strip()
            print("      caption: %s" % (caption[:110] or "(vacio)"))
            if alt:
                print("      alt    : %s" % alt[:110])
                print("               ^ si el caption se parece a esto, el parser")
                print("                 esta leyendo el alt-text. PARAR.")

    print("")
    print("3 · RATIO DE ESPAÑOL")
    if ratios:
        media = statistics.fmean(ratios)
        print("   filas con ratio: %d de %d" % (len(ratios), len(filas)))
        print("   media   : %.1f%%" % (media * 100))
        print("   mediana : %.1f%%" % (statistics.median(ratios) * 100))
        print("")
        if media <= 0.06:
            print("   ⚠⚠ DETENERSE. La media es %.1f%%, cerca del 3,29%% que daba" % (media * 100))
            print("      el analisis por alt-text. Antes de lanzar sobre los 1.000:")
            print("      abre tres perfiles en el navegador y compara sus captions")
            print("      con lo que hay en ig_raw/.")
        else:
            print("   La media NO esta cerca del 3%%, asi que el detector esta")
            print("   viendo texto real. Igual conviene abrir dos perfiles y")
            print("   comparar contra la pantalla.")
    else:
        print("   ninguna fila tiene ratio. Revisa el estado de los perfiles.")

    print("")
    print("4 · S6 y S8 · las dos categorias que estaban vacias")
    con_lender = [f for _, _, f in filas if f.get("menciona_lender")]
    con_preg = [f for _, _, f in filas if (f.get("comentarios_pregunta_calificacion") or 0) > 0]
    print("   con cuenta hipotecaria etiquetada: %d de %d" % (len(con_lender), len(filas)))
    for f in con_lender[:5]:
        print("      @%s -> %s" % (f.get("handle"), f.get("cuentas_hipotecarias_etiquetadas")))
    print("   con preguntas de calificacion en comentarios: %d de %d"
          % (len(con_preg), len(filas)))
    for f in con_preg[:5]:
        print("      @%s -> %d preguntas"
              % (f.get("handle"), f.get("comentarios_pregunta_calificacion")))

    return {
        "perfiles": len(filas),
        "por_estado": estados,
        "captions_es_ratio_media": statistics.fmean(ratios) if ratios else None,
        "sospecha_alt_text": bool(ratios) and statistics.fmean(ratios) <= 0.06,
    }


# ── CLI ───────────────────────────────────────────────────────────────────────

def _main(argv: list[str] | None = None) -> int:
    import argparse

    ap = argparse.ArgumentParser(
        description="Capa de Instagram: sesion, lote crudo, parser y CSV.",
    )
    ap.add_argument("--iniciar-sesion", action="store_true",
                    help="abre el navegador para iniciar sesion una vez")
    ap.add_argument("--objetivo", action="store_true",
                    help="imprime la composicion de la lista y no raspa nada")
    ap.add_argument("--lote", action="store_true",
                    help="pasada 1: raspa y guarda el crudo")
    ap.add_argument("--parsear", action="store_true",
                    help="pasada 2: crudo -> ig_signals.csv")
    ap.add_argument("--piloto", type=int, metavar="N", default=None,
                    help="pasada 1 sobre N perfiles (el brief pide 20)")
    ap.add_argument("--revisar-piloto", action="store_true",
                    help="imprime la revision a mano del piloto")
    ap.add_argument("--limite", type=int, default=None)
    ap.add_argument("--desde-cero", action="store_true",
                    help="ignora el checkpoint")
    ap.add_argument("--excluir-descartados", action="store_true",
                    help="saca del objetivo a los DESCARTADO y RECLASIFICADO")
    ap.add_argument("--sin-comentarios", action="store_true")
    ap.add_argument("--con-ventana", action="store_true",
                    help="no headless: obligatorio en el piloto")
    args = ap.parse_args(argv)

    if args.iniciar_sesion:
        try:
            try:
                from patchright.sync_api import sync_playwright
            except ImportError:
                from playwright.sync_api import sync_playwright
        except ImportError as exc:
            raise SystemExit(
                "Falta el navegador:\n"
                "    pip install -r realtor_scraper/requirements.txt\n"
                "    python -m playwright install chromium\n"
                "\nDetalle: %s" % exc
            ) from exc
        from navegador import iniciar_sesion_instagram

        with sync_playwright() as p:
            return 0 if iniciar_sesion_instagram(p) else 1

    if args.objetivo:
        _, informe = cargar_objetivo(excluir_descartados=args.excluir_descartados)
        print(json.dumps(informe, ensure_ascii=False, indent=2, default=str))
        return 0

    if args.revisar_piloto:
        revisar_piloto()
        return 0

    if args.piloto is not None or args.lote:
        n = args.piloto if args.piloto is not None else args.limite
        resumen = correr_lote(
            limite=n,
            reanudar=not args.desde_cero,
            excluir_descartados=args.excluir_descartados,
            headless=not args.con_ventana,
            con_comentarios=not args.sin_comentarios,
        )
        print(json.dumps(resumen, ensure_ascii=False, indent=2, default=str))
        if args.piloto is not None:
            print("")
            print("Ahora la revision a mano, que el brief pide ANTES de los 1.000:")
            print("    python -m instagram.finder --revisar-piloto")
        return 0 if not resumen.get("detenido_por") else 1

    if args.parsear:
        informe = parsear_crudos()
        print(json.dumps(informe, ensure_ascii=False, indent=2, default=str))
        return 0

    ap.print_help()
    return 2


if __name__ == "__main__":
    import sys

    sys.exit(_main())
