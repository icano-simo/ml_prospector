"""Toda la ficha de UN realtor con consultas que cuestan 0. Prueba de concepto.

Quita las columnas que venian de llamadas por fila --los listados de lenders y
originadores, y el LO principal-- y las reemplaza por lo que se puede saber
preguntando, que no cobra:

  · Everett: si/no y CUANTAS operaciones, en dos ventanas;
  · tipo de programa: FHA, convencional, VA;
  · historical units por año, de 2017 a hoy, mas el año en curso;
  · de ahi se deriva el primer año con produccion y los años produciendo.

La ficha del agente (1 credito) NO se vuelve a pedir: se lee del crudo que ya
esta en `data/raw/`. Asi la corrida entera cuesta 0 y lo demuestra leyendo el
ledger antes y despues.

Uso:
    python modelmatch/ficha_cero.py                  # el agente de prueba
    python modelmatch/ficha_cero.py mma_xxxxxxxx     # uno concreto
"""
from __future__ import annotations

import glob
import json
import os
import sys

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from modelmatch.auditar_casa import CONFIRMADOS, DUDOSOS  # noqa: E402
from modelmatch.cliente import llamar, saldo, un_agente_de_prueba  # noqa: E402

EVERETT = CONFIRMADOS + DUDOSOS
SALIDA = os.path.join(RAIZ, "data", "salida")
CRUDO = os.path.join(RAIZ, "data", "raw")

#: Años COMPLETOS. El año en curso NO se pide como año literal: se comprobo
#: que `period: "2026"` devuelve 0 para un agente que en `yearToDate` tiene
#: mas de cinco operaciones. El bucket del año corriente no esta poblado, y
#: pedirlo asi hace creer que el realtor dejo de producir. Para el año en
#: curso va `yearToDate`.
import datetime as _dt  # noqa: E402

ANIO_EN_CURSO = _dt.date.today().year
ANIOS = [str(a) for a in range(2017, ANIO_EN_CURSO)]
#: Bandas para acorralar un numero de unidades. Finas abajo, gruesas arriba:
#: la diferencia entre 2 y 3 operaciones decide, la de 60 a 70 no.
BANDAS = [(1, 2), (2, 3), (3, 4), (4, 5), (5, 7), (7, 10), (10, 15),
          (15, 20), (20, 30), (30, 50), (50, 100), (100, None)]
#: Tipos de programa. El si/no es seguro; las UNIDADES por programa no, porque
#: se miden dentro del bucket de un lender suelto.
PROGRAMAS = (("fha", "FHA"), ("conventional", "convencional"), ("va", "VA"))


def contar(flat: dict, periodo: str, pie: dict | None = None,
           etq: str = "") -> int | None:
    cuerpo: dict = {"flatFilters": flat, "period": periodo}
    if pie:
        cuerpo["footprint"] = pie
    d = llamar("/v1/agents/count", cuerpo, etiqueta=etq, silencioso=True)
    if not isinstance(d, dict) or not isinstance(d.get("total"), (int, float)):
        return None
    return int(d["total"])


def acorralar(mm: str, periodo: str, campo: str = "buyerUnits",
              pie_base: dict | None = None) -> object:
    """El numero exacto, o una banda, o '' si no hubo nada. Cuesta 0."""
    def hay(gte, lt):
        rango: dict = {"gte": gte}
        if lt is not None:
            rango["lt"] = lt
        if pie_base is not None:
            pie = dict(pie_base)
            pie["units"] = rango
            return contar({"id": mm}, periodo, pie, "q_%s_%s" % (periodo, gte))
        return contar({"id": mm, campo: rango}, periodo, None,
                      "q_%s_%s" % (periodo, gte))

    if (hay(1, None) or 0) < 1:
        return ""
    for gte, lt in BANDAS:
        if hay(gte, lt) == 1:
            if lt is None:
                return "%d+" % gte
            return gte if lt == gte + 1 else "%d-%d" % (gte, lt - 1)
    return "100+"


def ficha_guardada(mm: str) -> dict:
    """La ficha ya pagada, del crudo. NO se vuelve a pedir."""
    for patron in ("mm_det_%s_*.json" % mm, "mm_agent_detalle_*.json"):
        for a in sorted(glob.glob(os.path.join(CRUDO, patron)), reverse=True):
            with open(a, encoding="utf-8") as fh:
                d = json.load(fh)
            data = d.get("data")
            if isinstance(data, dict) and data.get("id") == mm:
                return data
    return {}


def main() -> None:
    mm = next((a for a in sys.argv[1:] if a.startswith("mma_")), None) \
        or un_agente_de_prueba()
    s0 = saldo()
    print("agente: %s" % mm)
    print("saldo al empezar: %s · esto NO debe gastar nada" % s0)
    print("")

    det = ficha_guardada(mm)
    if not det:
        raise SystemExit("no tengo su ficha guardada en data/raw; este script "
                         "no la compra, para no gastar")
    s = det.get("scored") or {}
    print("ficha leida del crudo ya pagado · %s" % det.get("fullName"))

    filas: list[tuple] = []

    def add(grupo, campo, valor, origen, costo):
        filas.append((grupo, campo, valor, origen, costo))

    # ── de la ficha ya pagada ──────────────────────────────────────────────
    for campo, valor in (
            ("ID Model Match", det.get("id")),
            ("Brokerage (Model Match, hoy)", det.get("office")),
            ("Ciudad (MM)", det.get("city")),
            ("Estado (MM)", det.get("state")),
            ("ZIP (MM)", det.get("zip")),
            ("Licencia (MM)", det.get("licenseNumber")),
            ("Unidades 12m (MM)", det.get("units")),
            ("Volumen 12m (MM)", det.get("volume")),
            ("Precio medio", det.get("avgSoldPrice")),
            ("Compras (u)", det.get("buyerUnits")),
            ("Compras ($)", det.get("buyerVolume")),
            ("Ventas (u)", det.get("sellerUnits")),
            ("Ventas ($)", det.get("sellerVolume")),
            ("Dual (u)", det.get("dualUnits")),
            ("Compras FINANCIADAS (u)",
             s.get("total_mortgaged_buyer_units")),
            ("% unidades financiadas",
             s.get("total_percent_units_mortgaged")),
            ("Loan medio de sus compradores",
             s.get("average_mortgaged_loan_amount")),
            ("Nº lenders", det.get("totalLendersWorkedWith")),
            ("Nº originadores", det.get("totalOriginatorsWorkedWith")),
            ("Nº compañías", det.get("totalCompaniesWorkedWith"))):
        add("Ficha (ya pagada)", campo, valor, "GET /v1/agents/{id}", "1")

    # ── Everett: dos ventanas, con el numero ───────────────────────────────
    print("")
    print("── Everett ──")
    for periodo, etiqueta in (("allTime", "histórico"),
                              ("last12Months", "12 meses")):
        n = acorralar(mm, periodo, pie_base={"lender": EVERETT})
        add("Everett", "¿Trabajó con Everett · %s?" % etiqueta,
            "SÍ" if n != "" else "no", "count + footprint.lender", "0")
        add("Everett", "Operaciones con Everett · %s" % etiqueta, n,
            "count + footprint.lender + units", "0")
        print("   %-10s %s" % (etiqueta, n if n != "" else "sin operaciones"))

    # ── programas ──────────────────────────────────────────────────────────
    print("")
    print("── programas ──")
    for clave, nombre in PROGRAMAS:
        t = contar({"id": mm}, "last24Months",
                   {"dimension": "lender",
                    "product": {"mix": "loanType", "key": clave}},
                   "prog_%s" % clave)
        v = "sin comprobar" if t is None else ("sí" if t >= 1 else "no")
        add("Programas", "¿Produce %s?" % nombre, v,
            "count + footprint.product", "0")
        print("   %-14s %s" % (nombre, v))

    # ── historical units por año ───────────────────────────────────────────
    print("")
    print("── historical units · buy side por año ──")
    primer = None
    anios_con = 0
    for a in ANIOS:
        n = acorralar(mm, a)
        add("Historical units", "Historical units · %s" % a, n,
            "count + buyerUnits por año", "0")
        if n != "":
            primer = primer or a
            anios_con += 1
            print("   %s  %s" % (a, n))

    # El año en curso y el trimestre movil. NO hay trimestres calendario en
    # el `period` de agentes: `last3Months` es movil, no Q4.
    for periodo, etiqueta in (("yearToDate", "año en curso"),
                              ("last3Months", "últimos 3 meses")):
        n = acorralar(mm, periodo)
        add("Historical units", "Historical units · %s" % etiqueta, n,
            "count + buyerUnits", "0")
        print("   %-16s %s" % (etiqueta, n))

    # ── derivados, sin llamadas ────────────────────────────────────────────
    add("Derivado", "Primer año con producción", primer or "",
        "derivado de los años", "0")
    add("Derivado", "Años con producción", anios_con, "derivado", "0")
    # Antigüedad APROXIMADA, y hay que decir por que. El primer año con
    # produccion es un piso, no la fecha en que empezo: si su primera
    # operacion es anterior a 2017 --el año mas viejo que acepta el enum--
    # este numero la subestima y no hay forma de saberlo desde aqui.
    add("Derivado", "Antigüedad aproximada (años)",
        (ANIO_EN_CURSO - int(primer) + 1) if primer else "",
        "derivado del primer año", "0")
    add("Derivado", "¿La antigüedad está topada por 2017?",
        "sí · puede ser mayor" if primer == "2017" else "no", "derivado", "0")

    # ── el Excel ───────────────────────────────────────────────────────────
    wb = Workbook()
    ws = wb.active
    ws.title = "Ficha a coste cero"
    ws.append(["Grupo", "Campo", "Valor", "De dónde sale", "Créditos"])
    for f in filas:
        ws.append(list(f))
    for i, an in enumerate([22, 38, 26, 34, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = an
        c = ws.cell(row=1, column=i)
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    verde = PatternFill("solid", fgColor="E2EFDA")
    for i in range(2, len(filas) + 2):
        if ws.cell(row=i, column=5).value == "0":
            ws.cell(row=i, column=5).fill = verde
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    os.makedirs(SALIDA, exist_ok=True)
    ruta = os.path.join(SALIDA, "ficha_coste_cero.xlsx")
    try:
        wb.save(ruta)
    except PermissionError:
        import datetime as dt
        ruta = os.path.join(SALIDA, "ficha_coste_cero_%s.xlsx"
                            % dt.datetime.now().strftime("%H%M"))
        wb.save(ruta)

    s1 = saldo()
    gasto = None if None in (s0, s1) else round(s0 - s1, 2)
    print("")
    print("campos en el Excel: %d" % len(filas))
    print("   de la ficha ya pagada: %d"
          % sum(1 for f in filas if f[4] == "1"))
    print("   nuevos, a coste cero : %d"
          % sum(1 for f in filas if f[4] == "0"))
    print("")
    print("saldo: %s · gasto de esta corrida: %s  %s"
          % (s1, gasto, "OK, fue GRATIS" if gasto == 0 else "⚠ COBRO"))
    print("guardado en %s" % ruta)


if __name__ == "__main__":
    main()
